"""Tests for the provenance marker helper."""

from __future__ import annotations

import json

import pytest

from scripts import provenance


def test_text_marker_round_trips():
    stamped = provenance.stamp_text("hello world\n")
    assert provenance.read_marker(stamped) == provenance.WORD


def test_text_marker_adds_no_word_and_no_visible_whitespace():
    stamped = provenance.stamp_text("the quick brown fox\n")
    # The marker is zero-width, so it neither adds a token nor a visible glyph.
    assert len(stamped.split()) == 4
    visible = "".join(ch for ch in stamped if ch not in provenance._ZERO_WIDTH_CHARS)
    assert visible.strip() == "the quick brown fox"


def test_text_marker_keeps_visible_content_byte_for_byte():
    original = "First line.\nSecond line.\n"
    stamped = provenance.stamp_text(original)
    visible = "".join(ch for ch in stamped if ch not in provenance._ZERO_WIDTH_CHARS)
    assert visible == original


def test_text_stamp_is_idempotent():
    once = provenance.stamp_text("essay body\n")
    twice = provenance.stamp_text(once)
    assert once == twice
    assert provenance.read_marker(twice) == provenance.WORD


def test_json_marker_round_trips_and_stays_valid_json():
    payload = json.dumps({"score": 90, "note": "good"}, indent=2)
    stamped = provenance.stamp_json(payload)
    assert json.loads(stamped) == {"score": 90, "note": "good"}
    assert provenance.read_marker(stamped) == provenance.WORD


def test_json_stamp_is_idempotent():
    once = provenance.stamp_json('{"a": 1}')
    twice = provenance.stamp_json(once)
    assert once == twice
    assert provenance.read_marker(twice) == provenance.WORD


def test_csv_marker_does_not_add_a_data_row():
    import csv
    import io

    original = "name,student_id,email\nAnna Chen,1,anna@example.edu\n"
    stamped = provenance.stamp_csv(original)
    rows = [
        row
        for row in csv.DictReader(io.StringIO(stamped, newline=""))
        if (row.get("name") or "").strip()
    ]
    assert [row["name"] for row in rows] == ["Anna Chen"]
    assert provenance.read_marker(stamped) == provenance.WORD


def test_csv_stamp_is_idempotent():
    original = "name\nAnna Chen\n"
    once = provenance.stamp_csv(original)
    twice = provenance.stamp_csv(once)
    assert once == twice


def test_stamp_bytes_dispatches_by_suffix():
    data = provenance.stamp_bytes(b"body text\n", ".md")
    assert provenance.read_marker(data.decode("utf-8")) == provenance.WORD


def test_read_marker_absent_returns_none():
    assert provenance.read_marker("plain text with no marker\n") is None


def test_docx_marker_lives_in_keywords_and_not_in_surfaced_metadata(tmp_path):
    import docx

    from scripts.redaction.documents import _docx_metadata

    document = docx.Document()
    document.add_paragraph("Body of the essay.")
    document.core_properties.author = "Anna Chen"
    source = tmp_path / "essay.docx"
    document.save(source)

    stamped = provenance.stamp_bytes(source.read_bytes(), ".docx")
    source.write_bytes(stamped)

    reopened = docx.Document(str(source))
    assert provenance.read_marker(reopened.core_properties.keywords) == provenance.WORD
    # Keywords must never be surfaced as identity-bearing metadata.
    assert all("keyword" not in value.lower() for value in _docx_metadata(reopened))
    assert reopened.core_properties.keywords not in _docx_metadata(reopened)


def test_docx_stamp_is_idempotent(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Body.")
    source = tmp_path / "essay.docx"
    document.save(source)

    once = provenance.stamp_bytes(source.read_bytes(), ".docx")
    twice = provenance.stamp_bytes(once, ".docx")
    reopened_once = docx.Document(__import__("io").BytesIO(once))
    reopened_twice = docx.Document(__import__("io").BytesIO(twice))
    assert (
        reopened_once.core_properties.keywords
        == reopened_twice.core_properties.keywords
    )


def test_strip_marker_recovers_original_text_bytes():
    original = b"Essay body line one.\nLine two.\n"
    stamped = provenance.stamp_bytes(original, ".md")
    assert stamped != original
    assert provenance.strip_marker(stamped, ".md") == original


def test_strip_marker_recovers_original_json_and_csv():
    for suffix, original in ((".json", b'{"a": 1}\n'), (".csv", b"name\nAnna\n")):
        stamped = provenance.stamp_bytes(original, suffix)
        assert provenance.strip_marker(stamped, suffix) == original
